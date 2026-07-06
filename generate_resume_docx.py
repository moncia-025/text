from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = "/workspace/陈子恒_AI产品_AIWorkflow_简历.docx"


def set_run_font(run, size=9, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph_spacing(paragraph, before=0, after=2, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_bottom_border(paragraph, color="2F5597", size="6"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def add_section(doc, title):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=3)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=(47, 85, 151))
    add_bottom_border(p)


def add_text(doc, text, size=9, bold=False, after=2):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style=None)
    set_paragraph_spacing(p, after=1.5, line=1.0)
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.28)
    r = p.add_run("• " + text)
    set_run_font(r, size=8.6)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, after=1)
    r = p.add_run("陈子恒")
    set_run_font(r, size=18, bold=True)

    add_text(
        doc,
        "151-3734-2002 ｜ xigua050205@163.com ｜ GitHub: github.com/moncia-25 ｜ 期望城市：不限 ｜ 到岗：随时",
        size=8.8,
        after=4,
    )

    add_section(doc, "教育背景")
    add_text(doc, "新乡工程学院｜计算机科学与技术｜本科在读（2023.09 - 2027.06）", size=9, bold=True)
    add_text(doc, "GPA 3.6/4.0｜英语四级｜国家励志奖学金｜相关课程：软件工程、数据库、计算机网络、Web 开发、数据结构", size=8.6)

    add_section(doc, "专业技能")
    skill_lines = [
        "AI 工具 / Agent：Dify Workflow、RAG 知识库、AI Agent 流程设计、Prompt Engineering、JSON 结构化输出、LLM 应用设计、Cursor 辅助开发、MCP 了解",
        "产品能力：PRD、用户场景拆解、业务流程图、字段设计、指标体系、异常流程设计、产品 Demo 验证、飞书文档",
        "数据分析：SQL 基础查询、Python 基础数据处理、Excel 透视表、数据看板、分类分布分析、人工介入率分析、指标拆解",
        "自动化工具：影刀 RPA、Dify API 调用、Excel 自动读取与回写、自动化流程设计",
        "前端 / 工程基础：Vue3、Next.js、React 了解、TypeScript 了解、Axios、ECharts、SSE 流式响应、GitHub、Vercel 部署",
    ]
    for line in skill_lines:
        add_text(doc, line, size=8.6, after=1.5)

    add_section(doc, "项目经历")
    add_text(doc, "TrendScout AI｜跨境电商 AI 趋势分析与选品辅助系统", size=9.4, bold=True)
    add_text(doc, "Demo: trendscout-ai-pied.vercel.app ｜ GitHub: github.com/moncia-25/trendscout-ai", size=8.3)
    trend_bullets = [
        "围绕跨境卖家“选品依赖经验、趋势判断缺少量化标准、进入时机滞后”的问题，设计 AI 趋势分析产品方案，拆解用户输入、趋势识别、产品评分、优先级排序和报告输出流程。",
        "基于 Dify Workflow 设计“需求输入 → 趋势分析 → 产品评分 → 选品建议 → 报告生成”AI 工作流，结合 Prompt 调优与结构化输出提升趋势分析结果的可读性和稳定性。",
        "设计 Trend Score 评分逻辑，从增长速度、社交热度、市场扩散度、生命周期阶段等维度辅助判断产品机会，用于支持选品优先级排序。",
        "使用 Cursor 辅助完成前端 Demo 搭建，并通过 Vercel 部署 Web 可视化页面，展示趋势榜单、产品分析结果和 AI 报告输出效果。",
        "通过 GitHub 管理项目代码与迭代过程，沉淀跨境电商 AI 选品场景下的需求分析、Workflow 设计和 Demo 验证经验。",
    ]
    for item in trend_bullets:
        add_bullet(doc, item)

    add_text(doc, "企业客服工单自动分派与回复助手｜Dify + RAG + 影刀 RPA + Excel 数据分析", size=9.4, bold=True, after=1)
    service_bullets = [
        "针对客服工单分类耗时、回复口径不统一和高优先级投诉易遗漏的问题，使用飞书文档完成项目 PRD，梳理目标用户、业务流程、输入输出字段、分类规则、优先级规则、异常处理和验收标准。",
        "基于 Dify Workflow + RAG 知识库搭建工单处理流程，实现工单分类、知识库检索、优先级判断、人工介入识别、客服回复生成与结构化 JSON 输出。",
        "使用影刀 RPA 读取 Excel 工单数据，调用 Dify Workflow API，并将分类、优先级、分派部门、回复建议和处理状态自动回写表格，跑通工单处理自动化闭环。",
        "将工单处理结果结构化为 category、priority、need_human、assigned_team、status 等字段，并基于 Excel 透视表分析工单分类分布、优先级占比、人工介入率和部门分派量。",
        "通过 Code 节点替代优先级判断、人工介入判断等确定性逻辑，降低 LLM 调用耗时并提升流程稳定性；基于数据分析结果沉淀知识库优化和客服策略迭代方向。",
    ]
    for item in service_bullets:
        add_bullet(doc, item)

    add_section(doc, "岗位匹配优势")
    advantages = [
        "具备计算机专业背景，能理解基础技术实现逻辑，并将业务问题拆解为字段、流程、规则、指标和可验证 Demo。",
        "有 Dify Workflow、RAG 知识库、Prompt 调优、影刀 RPA 和 Excel 数据分析的完整实践，能独立完成 AI 产品原型和自动化流程验证。",
        "相比传统开发，更关注 AI 工具在具体业务场景中的落地价值，能够通过 PRD、流程图、数据表和 Demo 展示产品方案。",
        "关注 AI SaaS、跨境电商 AI、Agent Workflow 和自动化工具方向，适合 AI 产品助理、AI Workflow、AI SaaS 产品运营、DevRel 和 AI 增长运营相关岗位。",
    ]
    for item in advantages:
        add_bullet(doc, item)

    add_section(doc, "自我评价")
    add_text(
        doc,
        "计算机科学与技术本科在读，具备 AI 产品设计、Agent Workflow 搭建和 Demo 落地经验。熟悉 Dify Workflow、RAG 知识库、Prompt Engineering、Cursor 辅助开发、影刀 RPA、Excel 透视表，并掌握 SQL 基础查询与 Python 基础数据处理能力。能够将业务问题拆解为字段、流程、规则、指标和可验证 Demo，更关注 AI 工具如何解决具体业务场景。",
        size=8.6,
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
